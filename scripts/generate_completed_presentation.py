#!/usr/bin/env python3
"""Generate a completed AITDR presentation from the uploaded PDF.

The script preserves slides 1-13 as rendered images, then adds 14 editable
slides: 7 for entrepreneuriat and 7 for gestion de projet.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_PDF = Path("/home/ubuntu/.cursor/projects/workspace/uploads/AITDR_pptv1_3c62.pdf")
OUTPUT_PPTX = ROOT / "presentation" / "AITDR_completed_after_slide_13.pptx"
SLIDES_TO_KEEP = 13

NAVY = RGBColor(7, 70, 98)
NAVY_DARK = RGBColor(5, 49, 70)
BLUE = RGBColor(23, 111, 155)
TEAL = RGBColor(29, 167, 164)
MUTED = RGBColor(125, 151, 161)
LIGHT = RGBColor(241, 246, 248)
VERY_LIGHT = RGBColor(248, 251, 252)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(42, 160, 90)
ORANGE = RGBColor(230, 147, 45)
RED = RGBColor(190, 65, 65)


def add_textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    font_size: int = 22,
    color: RGBColor = NAVY,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    font_name: str = "Aptos",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    return box


def set_shape_fill(shape, color: RGBColor, line: RGBColor | None = None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line or color


def add_decor(slide, section: str, number: str):
    # Top-right navigation dots, matching the original deck style.
    for i in range(5):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(11.35 + i * 0.25), Inches(0.38), Inches(0.13), Inches(0.13)
        )
        set_shape_fill(dot, NAVY)

    # Bottom-left navigation dots.
    for i in range(5):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.55 + i * 0.22), Inches(6.95), Inches(0.12), Inches(0.12)
        )
        set_shape_fill(dot, NAVY)

    # Thin top and bottom lines.
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.72), Inches(3.8), Inches(0.03))
    set_shape_fill(top_line, MUTED)
    bottom_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(6.78), Inches(3.2), Inches(0.03))
    set_shape_fill(bottom_line, MUTED)

    add_textbox(slide, section, 0.55, 0.22, 2.7, 0.24, font_size=8, color=MUTED, bold=True)
    add_textbox(slide, number, 12.15, 7.05, 0.55, 0.22, font_size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_title(slide, title: str, subtitle: str | None = None):
    add_textbox(
        slide,
        title,
        0.62,
        0.55,
        10.2,
        0.55,
        font_size=25,
        color=NAVY,
        bold=True,
        italic=True,
        font_name="Georgia",
    )
    if subtitle:
        add_textbox(slide, subtitle, 0.65, 1.08, 9.0, 0.35, font_size=12, color=MUTED)


def add_bullet_box(
    slide,
    title: str,
    bullets: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill=VERY_LIGHT,
    accent=TEAL,
):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(box, fill, line=MUTED)
    box.line.width = Pt(1)

    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    set_shape_fill(accent_bar, accent)

    add_textbox(slide, title, x + 0.22, y + 0.18, w - 0.38, 0.32, font_size=13, color=NAVY, bold=True)
    tf = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.6), Inches(w - 0.45), Inches(h - 0.75)).text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(10.5)
        p.font.color.rgb = NAVY_DARK
        p.space_after = Pt(4)


def add_metric_card(slide, value: str, label: str, x: float, y: float, w: float, accent=BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.05))
    set_shape_fill(card, LIGHT, line=LIGHT)
    add_textbox(slide, value, x + 0.12, y + 0.15, w - 0.24, 0.34, font_size=19, color=accent, bold=True)
    add_textbox(slide, label, x + 0.12, y + 0.57, w - 0.24, 0.32, font_size=8.5, color=NAVY_DARK)


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float, *, header=True):
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if header and r == 0 else (VERY_LIGHT if r % 2 else WHITE)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(8.5 if len(rows) > 5 else 9.5)
                p.font.color.rgb = WHITE if header and r == 0 else NAVY_DARK
                p.font.bold = bool(header and r == 0)
    return table


def add_preserved_pdf_slides(prs: Presentation, pdf_path: Path, keep_count: int):
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        doc = fitz.open(pdf_path)
        if doc.page_count < keep_count:
            raise ValueError(f"PDF has only {doc.page_count} pages; cannot keep {keep_count}.")

        for idx in range(keep_count):
            page = doc[idx]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            img_path = tmp / f"slide_{idx + 1:02d}.png"
            pix.save(img_path)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_picture(str(img_path), 0, 0, width=prs.slide_width, height=prs.slide_height)


def add_entrepreneuriat_slides(prs: Presentation):
    data = [
        (
            "Entrepreneuriat 1/7",
            "Opportunité de marché",
            "Un besoin clair pour un SOC cloud accessible aux PME",
            [
                ("Contexte", ["Cyberattaques automatisées en hausse", "PME marocaines exposées sans SOC interne", "Solutions SIEM traditionnelles coûteuses et complexes"]),
                ("Problème client", ["Manque de visibilité sécurité", "Temps de détection trop long", "Réponse manuelle et non industrialisée"]),
                ("Opportunité", ["Solution Azure-native reproductible", "Coût maîtrisé via PAYG", "Déploiement rapide avec Terraform"]),
            ],
        ),
        (
            "Entrepreneuriat 2/7",
            "Segments de clientèle",
            "Les premiers utilisateurs ciblés par AITDR",
            [
                ("PME / startups", ["Budget limité", "Besoin d'une protection simple", "Peu ou pas d'équipe SOC dédiée"]),
                ("RSSI / analystes SOC", ["Réduction de l'alert fatigue", "Centralisation Sentinel", "Réponse automatisée par SOAR"]),
                ("MSSP / écoles", ["Solution réplicable", "Support pédagogique", "Base pour services managés"]),
            ],
        ),
        (
            "Entrepreneuriat 3/7",
            "Proposition de valeur",
            "Détecter, corréler et répondre rapidement aux menaces cloud",
            [
                ("Valeur technique", ["Pipeline complet : logs, SQL, Sentinel, SOAR, ML", "Détection d'anomalies par Isolation Forest / DBSCAN", "Architecture Hub-and-Spoke sécurisée"]),
                ("Valeur opérationnelle", ["MTTD visé : moins de 5 minutes", "MTTC visé : moins de 30 secondes", "Moins d'interventions manuelles"]),
                ("Valeur économique", ["Prototype sous budget de 100 USD", "Infrastructure optimisée", "Alternative accessible aux SIEM coûteux"]),
            ],
        ),
        (
            "Entrepreneuriat 4/7",
            "Business Model Canvas",
            "Synthèse du modèle d'affaires AITDR",
            "bmc",
        ),
        (
            "Entrepreneuriat 5/7",
            "Analyse concurrentielle",
            "Positionnement face aux solutions SIEM/SOAR existantes",
            "competition",
        ),
        (
            "Entrepreneuriat 6/7",
            "Modèle de revenus",
            "Des offres progressives pour accompagner la maturité client",
            "pricing",
        ),
        (
            "Entrepreneuriat 7/7",
            "Faisabilité financière",
            "Maîtriser les coûts pour rendre la cybersécurité accessible",
            "finance",
        ),
    ]

    for idx, (section, title, subtitle, content) in enumerate(data, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_decor(slide, section, f"{13 + idx}")
        add_title(slide, title, subtitle)
        if isinstance(content, list):
            x_positions = [0.72, 4.62, 8.52]
            accents = [BLUE, TEAL, ORANGE]
            for x, accent, (box_title, bullets) in zip(x_positions, accents, content):
                add_bullet_box(slide, box_title, bullets, x, 1.75, 3.35, 4.15, accent=accent)
            if idx == 3:
                add_metric_card(slide, "< 5 min", "Temps moyen de détection visé", 1.1, 6.05, 2.3, BLUE)
                add_metric_card(slide, "< 30 s", "Temps moyen de confinement visé", 5.35, 6.05, 2.3, TEAL)
                add_metric_card(slide, "100 USD", "Budget initial du prototype", 9.55, 6.05, 2.3, GREEN)
        elif content == "bmc":
            boxes = [
                ("Partenaires", "Azure, open source,\nécoles, TI feeds"),
                ("Activités", "Terraform, KQL,\nSOAR, ML, support"),
                ("Valeur", "SOC cloud-native\naccessible et automatisé"),
                ("Relations", "Documentation,\nonboarding, support"),
                ("Segments", "PME, RSSI,\nMSSP, écoles"),
                ("Ressources", "Sentinel, SQL,\nLogic Apps, scripts"),
                ("Canaux", "GitHub, Azure\nMarketplace, LinkedIn"),
                ("Coûts", "Cloud, maintenance,\nR&D, support"),
                ("Revenus", "Freemium, SaaS,\nEnterprise, formation"),
            ]
            for i, (box_title, body) in enumerate(boxes):
                col = i % 3
                row = i // 3
                add_bullet_box(slide, box_title, [body], 0.82 + col * 4.05, 1.55 + row * 1.65, 3.45, 1.25, accent=[BLUE, TEAL, ORANGE][col])
        elif content == "competition":
            rows = [
                ["Solution", "Forces", "Limites", "Position AITDR"],
                ["Microsoft Sentinel", "Azure natif, IA", "Configuration et coûts", "Simplification + playbooks prêts"],
                ["Splunk ES", "Écosystème mature", "Coût élevé", "Alternative PME abordable"],
                ["Elastic SIEM", "Flexible, open source", "Expertise requise", "Déploiement guidé Azure"],
                ["AITDR", "Vertical : capture à réponse", "Notoriété à construire", "Prototype reproductible"],
            ]
            add_table(slide, rows, 0.75, 1.6, 11.75, 4.6)
        elif content == "pricing":
            rows = [
                ["Offre", "Prix cible", "Fonctionnalités", "Cible"],
                ["Discovery", "Gratuit", "Open source, documentation", "Étudiants / chercheurs"],
                ["Basic", "1 200 MAD/mois", "Honeypot, Sentinel, règles KQL", "Micro-entreprises"],
                ["Premium", "2 500 MAD/mois", "SOAR, ML, Power BI, support", "PME / ETI"],
                ["Enterprise", "Sur devis", "SOC as a Service, formation", "MSSP / grands comptes"],
            ]
            add_table(slide, rows, 0.75, 1.6, 11.75, 4.55)
        elif content == "finance":
            add_metric_card(slide, "87,40 USD", "Coût final estimé du prototype", 0.9, 1.75, 2.7, GREEN)
            add_metric_card(slide, "11 clients", "Seuil de rentabilité estimé", 5.25, 1.75, 2.7, BLUE)
            add_metric_card(slide, "42%", "Marge nette cible S2", 9.55, 1.75, 2.7, TEAL)
            add_bullet_box(
                slide,
                "Hypothèses financières",
                [
                    "Revenu moyen annuel : 18 000 MAD/client",
                    "Charges fixes : support, maintenance, marketing",
                    "Coût variable principal : infrastructure Azure",
                    "Rentabilité atteinte lorsque la marge couvre les charges fixes",
                ],
                1.05,
                3.25,
                11.0,
                2.7,
                accent=GREEN,
            )


def add_project_management_slides(prs: Presentation):
    slides = [
        (
            "Gestion de projet 1/7",
            "Méthodologie retenue",
            "Une approche Agile/Scrum adaptée à un projet technique évolutif",
            [
                ("Pourquoi Agile ?", ["Validation incrémentale", "Adaptation aux contraintes Azure", "Réduction des risques par itérations"]),
                ("Organisation", ["Backlog fonctionnel", "Sprints courts orientés livrables", "Démonstrations régulières"]),
                ("Résultat", ["Apprentissage continu", "Priorisation par valeur", "Meilleure maîtrise du périmètre"]),
            ],
        ),
        (
            "Gestion de projet 2/7",
            "Découpage du travail",
            "Work Breakdown Structure du projet AITDR",
            [
                ("Infrastructure", ["VNet, NSG, Key Vault", "VM1/VM2, SQL, Storage", "Terraform modules"]),
                ("Sécurité & SIEM", ["Suricata, logs Linux", "Microsoft Sentinel", "Règles KQL"]),
                ("Automatisation & IA", ["Logic Apps SOAR", "Scripts Python ETL", "Isolation Forest / DBSCAN"]),
            ],
        ),
        (
            "Gestion de projet 3/7",
            "Planning par phases",
            "Progression logique depuis la conception jusqu'à la validation",
            "timeline",
        ),
        (
            "Gestion de projet 4/7",
            "Équipe et responsabilités",
            "Un projet réalisé par une équipe réduite avec rôles cumulés",
            "roles",
        ),
        (
            "Gestion de projet 5/7",
            "Gestion des risques",
            "Identifier les risques techniques et budgétaires dès la conception",
            "risks",
        ),
        (
            "Gestion de projet 6/7",
            "Indicateurs de pilotage",
            "Mesurer la réussite technique, opérationnelle et financière",
            "kpis",
        ),
        (
            "Gestion de projet 7/7",
            "Livrables et perspectives",
            "Ce que le projet livre et ce qui peut être industrialisé ensuite",
            "deliverables",
        ),
    ]

    for idx, (section, title, subtitle, content) in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_decor(slide, section, f"{20 + idx}")
        add_title(slide, title, subtitle)

        if isinstance(content, list):
            x_positions = [0.72, 4.62, 8.52]
            accents = [BLUE, TEAL, ORANGE]
            for x, accent, (box_title, bullets) in zip(x_positions, accents, content):
                add_bullet_box(slide, box_title, bullets, x, 1.75, 3.35, 4.15, accent=accent)
        elif content == "timeline":
            phases = [
                ("1", "Cadrage", "Besoin, objectifs,\narchitecture cible"),
                ("2", "Infrastructure", "Terraform, VNet,\nVM, SQL, Key Vault"),
                ("3", "Collecte", "DVWA, Suricata,\nETL Python"),
                ("4", "SIEM/SOAR", "Sentinel, KQL,\nLogic Apps"),
                ("5", "Validation", "ML, Power BI,\ncoûts, documentation"),
            ]
            y = 3.45
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(y + 0.18), Inches(10.9), Inches(0.04))
            set_shape_fill(line, MUTED)
            for i, (num, name, desc) in enumerate(phases):
                x = 1.05 + i * 2.35
                circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y - 0.12), Inches(0.46), Inches(0.46))
                set_shape_fill(circle, [BLUE, TEAL, ORANGE, GREEN, NAVY][i])
                add_textbox(slide, num, x + 0.15, y - 0.02, 0.17, 0.14, font_size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
                add_textbox(slide, name, x - 0.25, y + 0.55, 1.1, 0.25, font_size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
                add_textbox(slide, desc, x - 0.55, y + 0.95, 1.75, 0.72, font_size=8.5, color=NAVY_DARK, align=PP_ALIGN.CENTER)
        elif content == "roles":
            rows = [
                ["Rôle", "Responsabilités principales"],
                ["Chef de projet", "Planification, suivi, budget, risques, documentation"],
                ["Cloud Architect", "Architecture Azure, réseau, sécurité, Terraform"],
                ["Développeur SIEM/SOAR", "Sentinel, KQL, Logic Apps, automatisation"],
                ["Data Scientist", "Features, modèles ML, évaluation des anomalies"],
                ["DevSecOps", "Bootstrap Linux, Docker, systemd, scripts de collecte"],
            ]
            add_table(slide, rows, 1.0, 1.6, 11.25, 4.8)
        elif content == "risks":
            rows = [
                ["Risque", "Impact", "Mitigation"],
                ["Dépassement budget Azure", "Élevé", "Auto-shutdown, filtrage logs, suivi Cost Management"],
                ["Volume Sentinel trop élevé", "Élevé", "DCR ciblées, rétention maîtrisée, tests progressifs"],
                ["Données d'attaque insuffisantes", "Moyen", "Honeypot DVWA exposé et collecte continue"],
                ["Faux positifs ML", "Moyen", "Calibration, features simples, validation manuelle"],
                ["Complexité Terraform", "Moyen", "Modules séparés, plan avant apply, documentation"],
            ]
            add_table(slide, rows, 0.8, 1.55, 11.8, 4.95)
        elif content == "kpis":
            add_metric_card(slide, "< 1 h", "Déploiement infrastructure", 0.95, 1.65, 2.45, BLUE)
            add_metric_card(slide, "> 10k", "Requêtes malveillantes capturées", 3.95, 1.65, 2.45, TEAL)
            add_metric_card(slide, "< 5 min", "Détection Sentinel / pipeline", 6.95, 1.65, 2.45, ORANGE)
            add_metric_card(slide, "< 100 USD", "Respect du budget initial", 9.95, 1.65, 2.45, GREEN)
            add_bullet_box(
                slide,
                "Critères qualité",
                [
                    "Infrastructure reproductible via Terraform",
                    "Logs centralisés et requêtables dans Azure",
                    "Réponse SOAR testable et traçable",
                    "Tableaux de bord Power BI exploitables pour la décision",
                ],
                1.05,
                3.2,
                11.0,
                2.55,
                accent=NAVY,
            )
        elif content == "deliverables":
            add_bullet_box(
                slide,
                "Livrables réalisés",
                [
                    "Architecture Azure Hub-and-Spoke",
                    "Code Terraform modulaire",
                    "Honeypot DVWA et Suricata",
                    "Pipeline ETL Bash/Python vers SQL",
                    "Règles KQL, SOAR Logic Apps, ML et Power BI",
                ],
                0.85,
                1.55,
                5.4,
                4.45,
                accent=BLUE,
            )
            add_bullet_box(
                slide,
                "Perspectives",
                [
                    "Threat Intelligence et géolocalisation IP",
                    "Conversion Sigma vers KQL",
                    "Multi-honeypots géographiques",
                    "LLM pour rapports d'incident",
                    "Industrialisation SaaS / Marketplace Azure",
                ],
                7.05,
                1.55,
                5.4,
                4.45,
                accent=TEAL,
            )


def main():
    OUTPUT_PPTX.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    add_preserved_pdf_slides(prs, SOURCE_PDF, SLIDES_TO_KEEP)
    add_entrepreneuriat_slides(prs)
    add_project_management_slides(prs)

    prs.save(OUTPUT_PPTX)
    print(f"Generated {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
