#!/usr/bin/env python3
"""Generate a fully French AITDR presentation.

Outputs:
- presentation/AITDR_completed_after_slide_13_FR.pptx
- presentation/AITDR_completed_after_slide_13_FR.pdf
"""

from __future__ import annotations

import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, JpegImagePlugin  # noqa: F401
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "presentation"
OUT_PPTX = OUT_DIR / "AITDR_completed_after_slide_13_FR.pptx"
OUT_PDF = OUT_DIR / "AITDR_completed_after_slide_13_FR.pdf"
OUT_PDF_COMPAT = OUT_DIR / "AITDR_completed_after_slide_13.pdf"
OUT_NOTES = OUT_DIR / "AITDR_conception_speaker_notes.md"

SLIDE_W = 13.333333
SLIDE_H = 7.5

NAVY = RGBColor(7, 70, 98)
NAVY_DARK = RGBColor(5, 49, 70)
TEAL = RGBColor(26, 160, 158)
BLUE = RGBColor(24, 112, 160)
ORANGE = RGBColor(224, 142, 44)
GREEN = RGBColor(48, 156, 88)
RED = RGBColor(188, 72, 72)
MUTED = RGBColor(124, 150, 160)
LIGHT = RGBColor(242, 247, 249)
WHITE = RGBColor(255, 255, 255)

PIL_COLORS = {
    "navy": (7, 70, 98),
    "navy_dark": (5, 49, 70),
    "teal": (26, 160, 158),
    "blue": (24, 112, 160),
    "orange": (224, 142, 44),
    "green": (48, 156, 88),
    "red": (188, 72, 72),
    "muted": (124, 150, 160),
    "light": (242, 247, 249),
    "white": (255, 255, 255),
    "black": (30, 40, 45),
}


SLIDES = [
    {
        "type": "title",
        "section": "Accueil",
        "title": "AITDR",
        "subtitle": "Azure Intelligent Threat Detection & Response",
        "body": [
            "Plateforme SOC Cloud intelligente",
            "Détection automatisée, réponse SOAR et analyse par IA",
            "Projet de Fin d'Année - 2025 / 2026",
        ],
    },
    {
        "type": "sommaire",
        "section": "Plan",
        "title": "Sommaire",
        "subtitle": "Plan de la présentation",
        "items": [
            "Contexte, problématique et objectifs",
            "Conception et modélisation",
            "Infrastructure Azure et budget",
            "Collecte des logs, SIEM Sentinel et SOAR",
            "Machine Learning et visualisation Power BI",
            "Résultats et validation",
            "Entrepreneuriat : marché, BMC et modèle économique",
            "Gestion de projet : méthode, planning, risques et livrables",
            "Conclusion et perspectives",
        ],
    },
    {
        "type": "cards",
        "section": "Equipe",
        "title": "Equipe du projet",
        "subtitle": "Membres et responsabilités principales",
        "cards": [
            ("Oussama Gammal", ["Chef de projet", "Architecture Azure", "SIEM/SOAR, ML et documentation"], "blue"),
            ("Chbane Labib Jad", ["Contribution technique", "Validation fonctionnelle", "Support de conception"], "teal"),
            ("Encadrement", ["Suivi académique", "Validation des livrables", "Orientation méthodologique"], "orange"),
        ],
    },
    {
        "type": "text",
        "section": "Introduction",
        "title": "Introduction générale",
        "subtitle": "Pourquoi AITDR ?",
        "paragraph": (
            "Les cyberattaques automatisées augmentent fortement dans les environnements cloud. "
            "Les solutions de défense traditionnelles restent souvent coûteuses, complexes et basées sur des règles statiques. "
            "AITDR propose une plateforme SOC cloud-native permettant de détecter, analyser et répondre aux menaces de manière rapide et automatisée."
        ),
        "bullets": [
            "Surveillance continue des événements de sécurité.",
            "Centralisation des logs dans Azure.",
            "Réponse automatisée aux incidents critiques.",
            "Maîtrise du coût dans un budget initial limité.",
        ],
    },
    {
        "type": "text",
        "section": "Contexte",
        "title": "Contexte et problématique",
        "subtitle": "Un besoin de cybersécurité accessible",
        "paragraph": (
            "De nombreuses PME ne disposent pas d'un SOC interne et ne peuvent pas financer des solutions SIEM/SOAR professionnelles très coûteuses. "
            "Pourtant, elles sont exposées aux attaques automatisées, aux scans, aux tentatives SSH et aux vulnérabilités applicatives."
        ),
        "bullets": [
            "Manque de visibilité sur les attaques en temps réel.",
            "Temps de détection et de réponse trop élevé.",
            "Difficulté à déployer une architecture sécurisée et reproductible.",
            "Contraintes fortes de budget et de compétences.",
        ],
    },
    {
        "type": "cards",
        "section": "Objectifs",
        "title": "Objectifs du projet",
        "subtitle": "Les objectifs techniques et opérationnels d'AITDR",
        "cards": [
            ("Déployer", ["Infrastructure Azure segmentée", "Terraform et Infrastructure as Code", "Réseau Hub-and-Spoke"], "blue"),
            ("Détecter", ["Honeypot DVWA", "Suricata IDS", "Règles KQL Sentinel"], "teal"),
            ("Répondre", ["Logic Apps SOAR", "Blocage IP automatique", "Alertes et confinement rapide"], "orange"),
            ("Analyser", ["Pipeline ETL", "SQL Database", "Machine Learning et Power BI"], "green"),
        ],
    },
    {
        "type": "section",
        "section": "Conception",
        "title": "Conception et modélisation",
        "subtitle": "Sous-titre : de l'architecture Azure aux interactions fonctionnelles",
    },
    {
        "type": "architecture_detail",
        "section": "Conception 1/4",
        "title": "Architecture Azure AITDR",
        "subtitle": "Vue logique de la plateforme SOC cloud-native",
        "talk": (
            "Cette slide présente l'architecture globale de AITDR. On retrouve une zone DMZ exposée avec le honeypot DVWA, "
            "une zone Data pour stocker les logs, et une zone SIEM/SOAR pour analyser les événements et automatiser la réponse. "
            "L'objectif est de séparer les composants sensibles et de garder une architecture sécurisée et reproductible avec Terraform."
        ),
    },
    {
        "type": "class_diagram",
        "section": "Conception 2/4",
        "title": "Diagramme de classes AITDR",
        "subtitle": "Objets principaux de la plateforme",
        "talk": (
            "Ce diagramme de classes montre les principaux objets manipulés par la solution. L'AttackEvent représente l'événement collecté, "
            "le LogPipeline assure la collecte et le parsing, le SIEMAlert génère les alertes, et le SOARPlaybook exécute les réponses automatiques. "
            "Les modèles ML et le dashboard Power BI complètent la chaîne d'analyse et de visualisation."
        ),
    },
    {
        "type": "sequence_diagram",
        "section": "Conception 3/4",
        "title": "Diagramme de séquence",
        "subtitle": "Flux de détection et de réponse automatique",
        "talk": (
            "Ce diagramme explique le déroulement d'un incident. L'attaquant génère une requête vers le honeypot, Suricata produit des événements, "
            "le pipeline collecte et analyse les logs, puis Sentinel déclenche une alerte. Si l'incident est confirmé, Logic Apps bloque automatiquement "
            "l'adresse IP via le NSG et les tableaux de bord sont rafraîchis."
        ),
    },
    {
        "type": "usecase_diagram",
        "section": "Conception 4/4",
        "title": "Diagramme de cas d'utilisation",
        "subtitle": "Interactions entre attaquant, analyste SOC et administrateur",
        "talk": (
            "Ce diagramme présente les principaux utilisateurs du système. L'attaquant génère des événements, l'analyste SOC consulte les incidents "
            "et les tableaux de bord, tandis que l'administrateur configure l'infrastructure, les règles KQL et les alertes Sentinel. "
            "La plateforme relie donc la capture, la détection, la réponse et le reporting."
        ),
    },
    {
        "type": "diagram",
        "section": "Architecture",
        "title": "Architecture globale AITDR",
        "subtitle": "Topologie Hub-and-Spoke sur Microsoft Azure",
        "nodes": [
            ("On-premises", "Kali / Windows\nGénération de logs", "blue"),
            ("DMZ", "VM1 + DVWA\nSuricata IDS", "orange"),
            ("Data", "SQL Database\nBlob Storage", "teal"),
            ("SIEM/SOAR", "Sentinel\nLogic Apps + ML", "green"),
            ("Visualisation", "Power BI\nReporting", "blue"),
        ],
    },
    {
        "type": "text",
        "section": "Architecture",
        "title": "Partie on-premises",
        "subtitle": "Simulation et génération des événements",
        "paragraph": (
            "La partie on-premises représente l'environnement local utilisé pour simuler des activités normales et malveillantes. "
            "Kali Linux permet de générer des scénarios d'attaque, tandis que Windows Server joue un rôle de machine interne produisant des journaux."
        ),
        "bullets": [
            "Simulation d'attaques contrôlées.",
            "Génération d'événements de sécurité.",
            "Transmission sécurisée vers Azure via VPN.",
            "Validation du pipeline de supervision.",
        ],
    },
    {
        "type": "table",
        "section": "Infrastructure",
        "title": "Partie DMZ et ressources Azure",
        "subtitle": "Zone exposée et surveillée",
        "columns": ["Ressource", "Service", "Rôle", "Coût estimé"],
        "rows": [
            ["VNet DMZ", "Azure VNet", "Isolation réseau", "Gratuit"],
            ["NSG", "Network Security Group", "Ports 80/443 et SSH restreint", "Gratuit"],
            ["VM1 WebServer", "Azure VM B2s", "Honeypot DVWA", "25-40 USD/mois"],
            ["IP publique", "Azure Public IP", "Exposition contrôlée", "3-5 USD/mois"],
            ["Disque managé", "SSD Standard", "Stockage système", "6-8 USD/mois"],
        ],
    },
    {
        "type": "text",
        "section": "Données",
        "title": "Partie base de données",
        "subtitle": "Centralisation et corrélation des logs",
        "paragraph": (
            "Les données collectées depuis le honeypot, les systèmes Linux et Suricata sont normalisées puis stockées dans une base de données. "
            "Cette centralisation facilite l'analyse, la corrélation des événements et l'alimentation des tableaux de bord."
        ),
        "bullets": [
            "Tables WebAttacks, SSHAttacks, PacketLogs et MLAnomalies.",
            "Insertion automatisée par scripts Python.",
            "Base exploitable par Power BI et les modèles ML.",
            "Historisation des événements pour analyse post-incident.",
        ],
    },
    {
        "type": "diagram",
        "section": "Pipeline",
        "title": "Pipeline de collecte et traitement",
        "subtitle": "Processus ETL automatisé",
        "nodes": [
            ("Extract", "Logs Apache\nAuth.log\nEVE JSON", "blue"),
            ("Transform", "Parsing Python\nClassification\nNormalisation", "teal"),
            ("Load", "Azure SQL\nBlob Storage\nLog Analytics", "orange"),
            ("Analyze", "Sentinel\nML\nPower BI", "green"),
        ],
    },
    {
        "type": "text",
        "section": "SIEM",
        "title": "Microsoft Sentinel",
        "subtitle": "Corrélation et détection des incidents",
        "paragraph": (
            "Microsoft Sentinel joue le rôle de SIEM cloud. Il centralise les journaux, applique des règles KQL et génère des incidents lorsque des comportements suspects sont détectés."
        ),
        "bullets": [
            "Collecte des logs via Log Analytics Workspace.",
            "Règles KQL pour brute force SSH, SQLi, RCE et scans.",
            "Classification des alertes par sévérité.",
            "Déclenchement des playbooks SOAR.",
        ],
    },
    {
        "type": "cards",
        "section": "SOAR & IA",
        "title": "SOAR et Machine Learning",
        "subtitle": "Automatiser la réponse et détecter les anomalies",
        "cards": [
            ("SOAR", ["Logic Apps déclenchées par Sentinel", "Blocage automatique d'IP malveillantes", "Réduction du temps de confinement"], "orange"),
            ("Machine Learning", ["Isolation Forest", "DBSCAN", "Détection d'anomalies inconnues"], "teal"),
            ("Bénéfice", ["Moins d'actions manuelles", "Réponse plus rapide", "Amélioration continue"], "green"),
        ],
    },
    {
        "type": "text",
        "section": "Visualisation",
        "title": "Power BI et validation Azure",
        "subtitle": "Tableaux de bord et suivi opérationnel",
        "paragraph": (
            "Power BI permet de visualiser les attaques, les sources malveillantes, les types d'incidents et les anomalies détectées. "
            "Azure Cost Management et Defender for Cloud servent à valider le budget et la posture de sécurité."
        ),
        "bullets": [
            "Dashboard Power BI connecté à Azure SQL.",
            "Suivi des coûts et du crédit Azure for Students.",
            "Validation des logs dans Log Analytics.",
            "Recommandations de sécurité Defender for Cloud.",
        ],
    },
    {
        "type": "metrics",
        "section": "Résultats",
        "title": "Résultats clés",
        "subtitle": "Validation technique du prototype",
        "metrics": [
            ("100 USD", "Budget initial", "blue"),
            ("87,40 USD", "Coût final estimé", "green"),
            ("< 5 min", "Détection visée", "teal"),
            ("< 30 s", "Confinement SOAR visé", "orange"),
        ],
        "bullets": [
            "Infrastructure reproductible avec Terraform.",
            "Honeypot opérationnel pour capturer des attaques réelles.",
            "Logs centralisés, analysés et visualisés.",
            "Solution adaptée à un contexte académique et PME.",
        ],
    },
    {
        "type": "section",
        "section": "Entrepreneuriat",
        "title": "Partie Entrepreneuriat",
        "subtitle": "Transformer AITDR en solution commercialisable",
    },
    {
        "type": "cards",
        "section": "Entrepreneuriat 1/7",
        "title": "Opportunité de marché",
        "subtitle": "Un besoin réel pour un SOC accessible",
        "cards": [
            ("Constat", ["Hausse des attaques cloud", "PME sans SOC interne", "Pénurie de profils cyber"], "blue"),
            ("Problème", ["Solutions SIEM coûteuses", "Déploiement complexe", "Réponse souvent manuelle"], "orange"),
            ("Opportunité", ["SOC Azure prêt à déployer", "Automatisation SOAR", "Coût maîtrisé"], "green"),
        ],
    },
    {
        "type": "cards",
        "section": "Entrepreneuriat 2/7",
        "title": "Segments de clientèle",
        "subtitle": "Les utilisateurs et acheteurs potentiels",
        "cards": [
            ("PME / startups", ["Budget limité", "Besoin de visibilité", "Protection rapide"], "blue"),
            ("RSSI / SOC", ["Réduction de l'alert fatigue", "Centralisation SIEM", "Réponse automatisée"], "teal"),
            ("MSSP / écoles", ["Plateforme réplicable", "Support pédagogique", "Base pour services managés"], "orange"),
        ],
    },
    {
        "type": "cards",
        "section": "Entrepreneuriat 3/7",
        "title": "Proposition de valeur",
        "subtitle": "Ce que AITDR apporte au client",
        "cards": [
            ("Technique", ["Détection temps réel", "SIEM + SOAR + ML", "Déploiement Terraform"], "blue"),
            ("Opérationnelle", ["MTTD réduit", "Blocage automatique", "Moins d'interventions manuelles"], "green"),
            ("Économique", ["Prototype sous 100 USD", "Alternative aux SIEM coûteux", "Offres progressives"], "orange"),
        ],
    },
    {
        "type": "bmc",
        "section": "Entrepreneuriat 4/7",
        "title": "Business Model Canvas",
        "subtitle": "Le modèle d'affaires de AITDR",
    },
    {
        "type": "table",
        "section": "Entrepreneuriat 5/7",
        "title": "Analyse concurrentielle",
        "subtitle": "Positionnement face aux solutions existantes",
        "columns": ["Solution", "Forces", "Limites", "Position AITDR"],
        "rows": [
            ["Microsoft Sentinel", "Azure natif, puissant", "Configuration et coût", "Version guidée et orientée PME"],
            ["Splunk ES", "Écosystème mature", "Très coûteux", "Alternative plus accessible"],
            ["Elastic SIEM", "Flexible, open source", "Expertise requise", "Déploiement Azure simplifié"],
            ["IBM QRadar", "Réputation enterprise", "Coût élevé", "Approche légère et pédagogique"],
            ["AITDR", "Vertical et reproductible", "Notoriété à construire", "SOC complet à coût maîtrisé"],
        ],
    },
    {
        "type": "table",
        "section": "Entrepreneuriat 6/7",
        "title": "Modèle de revenus",
        "subtitle": "Des offres progressives selon la maturité client",
        "columns": ["Offre", "Prix cible", "Fonctionnalités", "Cible"],
        "rows": [
            ["Discovery", "Gratuit", "Open source, documentation", "Étudiants / chercheurs"],
            ["Basic", "1 200 MAD/mois", "Honeypot, Sentinel, règles KQL", "Micro-entreprises"],
            ["Premium", "2 500 MAD/mois", "SOAR, ML, Power BI, support", "PME / ETI"],
            ["Enterprise", "Sur devis", "Infrastructure dédiée, formation", "MSSP / grands comptes"],
        ],
    },
    {
        "type": "metrics",
        "section": "Entrepreneuriat 7/7",
        "title": "Faisabilité financière",
        "subtitle": "Maîtriser les coûts pour rendre la cybersécurité accessible",
        "metrics": [
            ("87,40 USD", "Coût final estimé du prototype", "green"),
            ("11 clients", "Seuil de rentabilité estimé", "blue"),
            ("18 000 MAD", "Revenu annuel moyen par client", "teal"),
            ("42 %", "Marge nette cible S2", "orange"),
        ],
        "bullets": [
            "Le coût principal vient de l'infrastructure Azure et de l'ingestion Sentinel.",
            "Le modèle SaaS permet des revenus récurrents.",
            "Le freemium facilite l'adoption et la visibilité du produit.",
        ],
    },
    {
        "type": "section",
        "section": "Gestion de projet",
        "title": "Partie Gestion de projet",
        "subtitle": "Planifier, piloter et valider AITDR",
    },
    {
        "type": "cards",
        "section": "Gestion de projet 1/7",
        "title": "Méthodologie retenue",
        "subtitle": "Approche Agile / Scrum",
        "cards": [
            ("Pourquoi Agile ?", ["Projet technique évolutif", "Tests progressifs", "Adaptation rapide"], "blue"),
            ("Organisation", ["Backlog fonctionnel", "Sprints orientés livrables", "Validation incrémentale"], "teal"),
            ("Bénéfice", ["Risques réduits", "Meilleure visibilité", "Amélioration continue"], "green"),
        ],
    },
    {
        "type": "cards",
        "section": "Gestion de projet 2/7",
        "title": "Découpage du travail",
        "subtitle": "Work Breakdown Structure simplifiée",
        "cards": [
            ("Infrastructure", ["VNet, NSG, VM", "SQL, Storage, Key Vault", "Modules Terraform"], "blue"),
            ("Sécurité", ["Honeypot DVWA", "Suricata", "Sentinel et KQL"], "orange"),
            ("Automatisation", ["Logic Apps SOAR", "Scripts ETL Python", "ML et Power BI"], "teal"),
        ],
    },
    {
        "type": "timeline",
        "section": "Gestion de projet 3/7",
        "title": "Planning par phases",
        "subtitle": "Progression logique du projet",
        "steps": [
            ("1", "Cadrage", "Besoin, marché,\nchoix techniques"),
            ("2", "Infrastructure", "Terraform,\nréseau, sécurité"),
            ("3", "Collecte", "DVWA,\nSuricata, ETL"),
            ("4", "SIEM/SOAR", "Sentinel,\nKQL, Logic Apps"),
            ("5", "Validation", "ML, Power BI,\ncoûts, rapport"),
        ],
    },
    {
        "type": "table",
        "section": "Gestion de projet 4/7",
        "title": "Équipe et responsabilités",
        "subtitle": "Rôles nécessaires pour réaliser AITDR",
        "columns": ["Rôle", "Responsabilités principales"],
        "rows": [
            ["Chef de projet", "Planification, suivi, budget, risques, documentation"],
            ["Cloud Architect", "Architecture Azure, réseau, sécurité, Terraform"],
            ["Développeur SIEM/SOAR", "Sentinel, KQL, Logic Apps, automatisation"],
            ["Data Scientist", "Feature engineering, modèles ML, évaluation"],
            ["DevSecOps", "Linux, Docker, systemd, scripts de collecte"],
        ],
    },
    {
        "type": "table",
        "section": "Gestion de projet 5/7",
        "title": "Gestion des risques",
        "subtitle": "Anticiper les blocages techniques et budgétaires",
        "columns": ["Risque", "Impact", "Mitigation"],
        "rows": [
            ["Dépassement budget Azure", "Élevé", "Cost Management, auto-shutdown, filtrage"],
            ["Volume Sentinel trop élevé", "Élevé", "DCR ciblées, rétention contrôlée"],
            ["Données insuffisantes", "Moyen", "Déploiement précoce du honeypot"],
            ["Faux positifs ML", "Moyen", "Calibration et validation manuelle"],
            ["Complexité Terraform", "Moyen", "Modules séparés et déploiement progressif"],
        ],
    },
    {
        "type": "metrics",
        "section": "Gestion de projet 6/7",
        "title": "Indicateurs de pilotage",
        "subtitle": "Mesurer la réussite technique et opérationnelle",
        "metrics": [
            ("< 1 h", "Déploiement infrastructure", "blue"),
            ("> 10k", "Requêtes malveillantes capturées", "teal"),
            ("< 5 min", "Temps de détection visé", "orange"),
            ("< 100 USD", "Respect du budget initial", "green"),
        ],
        "bullets": [
            "Logs centralisés et requêtables.",
            "Règles KQL fonctionnelles.",
            "Réponse SOAR testable.",
            "Dashboard Power BI exploitable.",
        ],
    },
    {
        "type": "cards",
        "section": "Gestion de projet 7/7",
        "title": "Livrables et perspectives",
        "subtitle": "Résultats produits et améliorations futures",
        "cards": [
            ("Livrables", ["Infrastructure Azure", "Code Terraform", "Scripts ETL", "Règles KQL"], "blue"),
            ("Validation", ["Sentinel", "SOAR Logic Apps", "ML anomalies", "Power BI"], "green"),
            ("Perspectives", ["Threat Intelligence", "Sigma vers KQL", "Multi-honeypots", "SaaS Azure Marketplace"], "teal"),
        ],
    },
    {
        "type": "conclusion",
        "section": "Conclusion",
        "title": "Conclusion générale",
        "subtitle": "AITDR : une plateforme SOC cloud-native complète",
        "bullets": [
            "AITDR démontre la faisabilité d'un SOC intelligent sur Azure avec un budget contrôlé.",
            "La solution combine infrastructure, collecte, SIEM, SOAR, Machine Learning et visualisation.",
            "La partie entrepreneuriale montre un potentiel de marché pour les PME et les MSSP.",
            "La gestion de projet prouve que la solution a été planifiée, suivie et validée de manière structurée.",
        ],
    },
    {
        "type": "thanks",
        "section": "Fin",
        "title": "Merci pour votre attention",
        "subtitle": "Questions / Réponses",
    },
]


def add_textbox(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 20,
    color: RGBColor = NAVY,
    bold: bool = False,
    italic: bool = False,
    align=PP_ALIGN.LEFT,
    font: str = "Aptos",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def fill_shape(shape, color: RGBColor, line: RGBColor | None = None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = line or color


def add_decor(slide, section: str, slide_no: int):
    for i in range(5):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(11.35 + i * 0.25), Inches(0.38), Inches(0.13), Inches(0.13)
        )
        fill_shape(dot, NAVY)
    for i in range(5):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(0.55 + i * 0.22), Inches(6.95), Inches(0.12), Inches(0.12)
        )
        fill_shape(dot, NAVY)

    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.72), Inches(3.6), Inches(0.03))
    bottom_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.0), Inches(6.78), Inches(3.2), Inches(0.03))
    fill_shape(top_line, MUTED)
    fill_shape(bottom_line, MUTED)
    add_textbox(slide, section, 0.55, 0.22, 4.2, 0.22, size=8, color=MUTED, bold=True)
    add_textbox(slide, str(slide_no), 12.05, 7.05, 0.55, 0.22, size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_title(slide, item: dict):
    add_textbox(
        slide,
        item["title"],
        0.62,
        0.55,
        10.5,
        0.55,
        size=25,
        color=NAVY,
        bold=True,
        italic=True,
        font="Georgia",
    )
    subtitle = item.get("subtitle")
    if subtitle:
        add_textbox(slide, subtitle, 0.65, 1.08, 10.2, 0.32, size=12, color=MUTED)


def add_card(slide, title: str, bullets: list[str], x: float, y: float, w: float, h: float, accent: RGBColor):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_shape(box, LIGHT, MUTED)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    fill_shape(bar, accent)
    add_textbox(slide, title, x + 0.22, y + 0.18, w - 0.36, 0.3, size=13, color=NAVY, bold=True)
    tf = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.65), Inches(w - 0.42), Inches(h - 0.8)).text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(10.3)
        p.font.color.rgb = NAVY_DARK
        p.space_after = Pt(5)


def add_cards_slide(slide, item: dict):
    cards = item["cards"]
    cols = 4 if len(cards) == 4 else 3
    x0 = 0.72
    gap = 0.32
    w = (11.95 - (cols - 1) * gap) / cols
    y = 1.72
    h = 4.15
    for idx, (title, bullets, color_name) in enumerate(cards):
        add_card(slide, title, bullets, x0 + idx * (w + gap), y, w, h, rgb(color_name))


def add_text_slide(slide, item: dict):
    add_textbox(slide, item["paragraph"], 0.85, 1.62, 11.6, 1.25, size=16, color=NAVY_DARK)
    add_card(slide, "Points clés", item["bullets"], 1.1, 3.25, 11.0, 2.6, TEAL)


def add_table_slide(slide, item: dict):
    rows = [item["columns"]] + item["rows"]
    table_shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.75), Inches(1.65), Inches(11.85), Inches(4.75))
    table = table_shape.table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else LIGHT)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Aptos"
                p.font.size = Pt(8.5)
                p.font.bold = r == 0
                p.font.color.rgb = WHITE if r == 0 else NAVY_DARK


def add_metrics_slide(slide, item: dict):
    metrics = item["metrics"]
    x0 = 0.75
    gap = 0.35
    w = (11.85 - (len(metrics) - 1) * gap) / len(metrics)
    for idx, (value, label, color_name) in enumerate(metrics):
        x = x0 + idx * (w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.65), Inches(w), Inches(1.2))
        fill_shape(card, LIGHT, LIGHT)
        add_textbox(slide, value, x + 0.12, 1.83, w - 0.24, 0.35, size=18, color=rgb(color_name), bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, label, x + 0.12, 2.28, w - 0.24, 0.35, size=8.5, color=NAVY_DARK, align=PP_ALIGN.CENTER)
    if item.get("bullets"):
        add_card(slide, "Interprétation", item["bullets"], 1.1, 3.35, 11.0, 2.35, TEAL)


def add_diagram_slide(slide, item: dict):
    nodes = item["nodes"]
    x0 = 0.75
    gap = 0.35
    w = (11.85 - (len(nodes) - 1) * gap) / len(nodes)
    y = 2.35
    for idx, (name, body, color_name) in enumerate(nodes):
        x = x0 + idx * (w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.55))
        fill_shape(box, LIGHT, rgb(color_name))
        add_textbox(slide, name, x + 0.1, y + 0.18, w - 0.2, 0.25, size=11, color=rgb(color_name), bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, body, x + 0.12, y + 0.58, w - 0.24, 0.65, size=8.8, color=NAVY_DARK, align=PP_ALIGN.CENTER)
        if idx < len(nodes) - 1:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w - 0.03), Inches(y + 0.55), Inches(0.42), Inches(0.35))
            fill_shape(arrow, MUTED)


def add_architecture_detail_slide(slide, item: dict):
    def component(title: str, body: str, x: float, y: float, w: float, h: float, color: RGBColor):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        fill_shape(box, color, color)
        add_textbox(slide, title, x + 0.08, y + 0.12, w - 0.16, 0.28, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, body, x + 0.1, y + 0.48, w - 0.2, h - 0.55, size=8.5, color=WHITE, align=PP_ALIGN.CENTER)

    def zone(title: str, x: float, y: float, w: float, h: float, color: RGBColor):
        outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        outer.fill.background()
        outer.line.color.rgb = color
        outer.line.width = Pt(1.5)
        add_textbox(slide, title, x + 0.12, y + 0.08, w - 0.24, 0.25, size=10, color=color, bold=True, align=PP_ALIGN.CENTER)

    component("Internet / Attaquant", "Trafic HTTP/SSH", 5.35, 1.35, 2.2, 0.6, RED)
    component("VPN Gateway", "IKEv2", 4.15, 2.15, 1.55, 0.55, NAVY_DARK)
    component("Azure Firewall", "Key Vault + contrôle", 6.0, 2.15, 1.9, 0.55, NAVY_DARK)

    zone("VNet1 - DMZ", 0.55, 3.05, 3.1, 2.45, BLUE)
    component("VM1 WebServer", "Ubuntu 22.04\nDVWA + Suricata IDS", 0.85, 3.55, 2.5, 0.72, BLUE)
    component("NSG-DMZ", "Allow 80/443\nDeny-All", 0.85, 4.38, 2.5, 0.58, ORANGE)
    component("Public IP", "Exposition VM1", 0.85, 5.05, 2.5, 0.4, RED)

    zone("VNet2 - Data", 4.05, 3.05, 3.1, 2.45, TEAL)
    component("Azure SQL Database", "attack_logs + anomalies", 4.35, 3.55, 2.5, 0.58, TEAL)
    component("Azure Blob Storage", "Logs bruts / rapports", 4.35, 4.23, 2.5, 0.58, TEAL)
    component("Azure Data Factory", "Pipeline ETL", 4.35, 4.91, 2.5, 0.58, TEAL)

    zone("VNet3 - SIEM/SOAR", 7.55, 3.05, 3.55, 2.45, NAVY)
    component("Microsoft Sentinel", "SIEM + KQL + AMA", 7.9, 3.48, 2.85, 0.58, NAVY)
    component("Azure Logic Apps", "SOAR + Playbooks", 7.9, 4.18, 2.85, 0.58, NAVY)
    component("VM2 ML Engine", "Isolation Forest + DBSCAN", 7.9, 4.88, 2.85, 0.58, NAVY)

    component("Log Analytics Workspace", "Azure Monitor + DCR", 4.5, 5.75, 2.4, 0.55, NAVY_DARK)
    component("Power BI", "Dashboards + DAX", 8.0, 5.75, 2.6, 0.55, ORANGE)
    add_textbox(slide, "Flux clés : HTTP/SSH -> logs -> SQL/Log Analytics -> Sentinel -> Logic Apps -> NSG", 1.0, 6.35, 11.2, 0.3, size=10, color=MUTED, align=PP_ALIGN.CENTER)


def add_class_diagram_slide(slide, item: dict):
    classes = [
        ("AttackEvent", ["event_id", "timestamp", "src_ip", "attack_type"], ["detect()", "classify()", "toJSON()"], RED, 0.65, 1.55),
        ("HoneypotVM", ["vm_id", "public_ip", "vnet", "dvwa_enabled"], ["captureTraffic()", "sendLogs()", "getStatus()"], BLUE, 0.65, 3.55),
        ("LogPipeline", ["pipeline_id", "blob_url", "schedule", "status"], ["collectLogs()", "parseLogs()", "storeSQL()"], NAVY, 3.7, 1.55),
        ("MLModel", ["model_type", "threshold", "features"], ["fit()", "predict()", "flagAnomaly()"], NAVY, 3.7, 3.55),
        ("SIEMAlert", ["alert_id", "rule_name", "severity", "kql_query"], ["evaluate()", "createIncident()", "triggerSOAR()"], TEAL, 6.75, 1.55),
        ("SOARPlaybook", ["playbook_id", "trigger", "actions", "status"], ["execute()", "blockIP()", "sendAlert()"], TEAL, 6.75, 3.55),
        ("NSGRule", ["rule_id", "priority", "src_ip", "action"], ["block()", "allow()", "delete()"], ORANGE, 9.8, 2.25),
        ("PowerBIDashboard", ["dataset", "refresh_schedule", "visuals"], ["refresh()", "exportReport()"], NAVY_DARK, 9.8, 4.45),
    ]

    def cls(name, attrs, methods, color, x, y):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.35), Inches(1.55))
        fill_shape(box, color, color)
        add_textbox(slide, name, x + 0.05, y + 0.08, 2.25, 0.22, size=10.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, "\n".join(f"- {a}" for a in attrs), x + 0.12, y + 0.38, 2.1, 0.52, size=6.8, color=WHITE)
        add_textbox(slide, "\n".join(f"+ {m}" for m in methods), x + 0.12, y + 0.96, 2.1, 0.45, size=6.8, color=WHITE)

    for klass in classes:
        cls(*klass)
    add_textbox(slide, "Relations principales : AttackEvent -> LogPipeline -> SIEMAlert -> SOARPlaybook -> NSGRule", 1.0, 6.35, 11.2, 0.3, size=10, color=MUTED, align=PP_ALIGN.CENTER)


def add_sequence_diagram_slide(slide, item: dict):
    actors = [
        ("Attaquant", RED),
        ("DVWA\nHoneypot", BLUE),
        ("Suricata\nIDS/IPS", BLUE),
        ("Pipeline\nETL / ML", NAVY),
        ("Sentinel\nSIEM", TEAL),
        ("Logic Apps\nSOAR", TEAL),
        ("NSG", ORANGE),
    ]
    x0 = 0.65
    gap = 1.72
    top = 1.65
    bottom = 6.15
    xs = []
    for idx, (name, color) in enumerate(actors):
        x = x0 + idx * gap
        xs.append(x + 0.55)
        head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top), Inches(1.25), Inches(0.5))
        fill_shape(head, color, color)
        add_textbox(slide, name, x + 0.05, top + 0.08, 1.15, 0.3, size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.58), Inches(top + 0.6), Inches(0.02), Inches(bottom - top - 0.55))
        fill_shape(line, LIGHT, LIGHT)

    steps = [
        (0, 1, "1. Attaque HTTP/SSH", 2.55, RED),
        (1, 2, "2. EVE JSON", 3.05, BLUE),
        (1, 3, "3. Logs Apache/auth", 3.45, BLUE),
        (3, 4, "4. Syslogs via AMA", 4.1, TEAL),
        (3, 4, "5. Anomalies ML", 4.55, TEAL),
        (4, 5, "6. Déclenche Logic App", 5.05, TEAL),
        (5, 6, "7. Bloquer IP NSG", 5.45, ORANGE),
        (6, 5, "8. Confirmation", 5.85, ORANGE),
    ]
    for src, dst, label, y, color in steps:
        x1, x2 = xs[src], xs[dst]
        arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(min(x1, x2)), Inches(y), Inches(abs(x2 - x1)), Inches(0.12))
        fill_shape(arrow, color)
        add_textbox(slide, label, min(x1, x2) + 0.08, y - 0.18, abs(x2 - x1) + 0.2, 0.2, size=7.5, color=NAVY_DARK)


def add_usecase_diagram_slide(slide, item: dict):
    add_textbox(slide, "Plateforme AITDR", 5.25, 1.45, 2.7, 0.3, size=13, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    boundary = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.8), Inches(1.7), Inches(7.4), Inches(4.45))
    boundary.fill.background()
    boundary.line.color.rgb = LIGHT

    def actor(label: str, x: float, y: float):
        head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.22), Inches(y), Inches(0.35), Inches(0.35))
        head.fill.background()
        head.line.color.rgb = MUTED
        body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.38), Inches(y + 0.35), Inches(0.03), Inches(0.7))
        fill_shape(body, MUTED)
        add_textbox(slide, label, x - 0.15, y + 1.08, 0.95, 0.25, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)

    def usecase(text: str, x: float, y: float, w: float, color: RGBColor):
        oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(0.48))
        fill_shape(oval, color, color)
        add_textbox(slide, text, x + 0.08, y + 0.14, w - 0.16, 0.2, size=8.8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)

    actor("Attaquant", 0.8, 2.1)
    actor("Analyste SOC", 0.7, 4.55)
    actor("Admin", 11.25, 3.35)
    usecase("Capturer des attaques", 3.35, 2.0, 3.2, BLUE)
    usecase("Collecter les logs", 3.35, 2.85, 3.2, BLUE)
    usecase("Détecter anomalies (ML)", 3.35, 3.7, 3.2, BLUE)
    usecase("Générer alertes SIEM", 3.35, 4.55, 3.2, TEAL)
    usecase("Réponse auto. SOAR", 3.35, 5.4, 3.2, TEAL)
    usecase("Visualiser tableau bord", 6.9, 2.35, 3.0, NAVY)
    usecase("Consulter incidents", 6.9, 3.25, 3.0, NAVY)
    usecase("Configurer règles KQL", 6.9, 4.15, 3.0, NAVY)
    usecase("Déployer infrastructure", 6.9, 5.05, 3.0, ORANGE)
    add_textbox(slide, "Honeypot / VM        SIEM / SOAR        Visualisation        Administration", 3.2, 6.32, 7.2, 0.25, size=9, color=MUTED, align=PP_ALIGN.CENTER)


def add_timeline_slide(slide, item: dict):
    y = 3.32
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(y + 0.21), Inches(10.9), Inches(0.04))
    fill_shape(line, MUTED)
    colors = [BLUE, TEAL, ORANGE, GREEN, NAVY]
    for idx, (num, name, desc) in enumerate(item["steps"]):
        x = 1.05 + idx * 2.35
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y - 0.12), Inches(0.5), Inches(0.5))
        fill_shape(circle, colors[idx])
        add_textbox(slide, num, x + 0.16, y + 0.01, 0.16, 0.13, size=9, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, name, x - 0.3, y + 0.6, 1.2, 0.25, size=12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, desc, x - 0.55, y + 1.0, 1.75, 0.75, size=8.5, color=NAVY_DARK, align=PP_ALIGN.CENTER)


def add_bmc_slide(slide):
    boxes = [
        ("Partenaires clés", "Microsoft Azure\nOpen source\nÉcoles / MSSP"),
        ("Activités clés", "Terraform\nKQL / SOAR\nML et support"),
        ("Proposition de valeur", "SOC cloud intelligent,\nautomatisé et économique"),
        ("Relations clients", "Documentation\nOnboarding\nSupport"),
        ("Segments clients", "PME, startups\nRSSI, MSSP\nÉcoles"),
        ("Ressources clés", "Azure, Sentinel\nSQL, Logic Apps\nScripts Python"),
        ("Canaux", "GitHub\nAzure Marketplace\nLinkedIn"),
        ("Structure de coûts", "Infrastructure Azure\nMaintenance\nSupport / R&D"),
        ("Sources de revenus", "Freemium\nSaaS\nEnterprise / formation"),
    ]
    for idx, (title, body) in enumerate(boxes):
        col = idx % 3
        row = idx // 3
        add_card(slide, title, [body], 0.72 + col * 4.05, 1.5 + row * 1.62, 3.55, 1.2, [BLUE, TEAL, ORANGE][col])


def add_sommaire_slide(slide, item: dict):
    left = item["items"][:5]
    right = item["items"][5:]
    add_card(slide, "Première partie", left, 0.9, 1.55, 5.65, 4.7, BLUE)
    add_card(slide, "Deuxième partie", right, 6.95, 1.55, 5.35, 4.7, TEAL)


def add_section_slide(slide, item: dict):
    add_textbox(slide, item["title"], 1.0, 2.55, 11.35, 0.7, size=34, color=NAVY, bold=True, italic=True, align=PP_ALIGN.CENTER, font="Georgia")
    add_textbox(slide, item["subtitle"], 1.5, 3.35, 10.35, 0.45, size=18, color=MUTED, align=PP_ALIGN.CENTER)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.65), Inches(4.15), Inches(4.0), Inches(0.06))
    fill_shape(line, TEAL)


def add_title_slide(slide, item: dict):
    add_textbox(slide, item["title"], 1.0, 2.05, 11.35, 0.85, size=48, color=NAVY, bold=True, italic=True, align=PP_ALIGN.CENTER, font="Georgia")
    add_textbox(slide, item["subtitle"], 1.0, 2.95, 11.35, 0.42, size=22, color=NAVY, align=PP_ALIGN.CENTER)
    y = 3.55
    for idx, line in enumerate(item["body"]):
        add_textbox(slide, line, 1.0, y + idx * 0.38, 11.35, 0.3, size=13, color=MUTED, align=PP_ALIGN.CENTER)


def add_conclusion_slide(slide, item: dict):
    add_card(slide, "Synthèse", item["bullets"], 1.05, 1.65, 11.2, 4.7, GREEN)


def add_thanks_slide(slide, item: dict):
    add_textbox(slide, item["title"], 1.0, 2.8, 11.35, 0.65, size=34, color=NAVY, bold=True, italic=True, align=PP_ALIGN.CENTER, font="Georgia")
    add_textbox(slide, item["subtitle"], 1.0, 3.55, 11.35, 0.45, size=22, color=TEAL, align=PP_ALIGN.CENTER)


def rgb(name: str) -> RGBColor:
    return {
        "navy": NAVY,
        "blue": BLUE,
        "teal": TEAL,
        "orange": ORANGE,
        "green": GREEN,
        "red": RED,
    }[name]


def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for idx, item in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if item["type"] not in {"title", "section", "thanks"}:
            add_decor(slide, item["section"], idx)
            add_title(slide, item)
        elif item["type"] == "title":
            add_decor(slide, item["section"], idx)
        else:
            add_decor(slide, item["section"], idx)

        kind = item["type"]
        if kind == "title":
            add_title_slide(slide, item)
        elif kind == "sommaire":
            add_sommaire_slide(slide, item)
        elif kind == "cards":
            add_cards_slide(slide, item)
        elif kind == "text":
            add_text_slide(slide, item)
        elif kind == "table":
            add_table_slide(slide, item)
        elif kind == "metrics":
            add_metrics_slide(slide, item)
        elif kind == "diagram":
            add_diagram_slide(slide, item)
        elif kind == "architecture_detail":
            add_architecture_detail_slide(slide, item)
        elif kind == "class_diagram":
            add_class_diagram_slide(slide, item)
        elif kind == "sequence_diagram":
            add_sequence_diagram_slide(slide, item)
        elif kind == "usecase_diagram":
            add_usecase_diagram_slide(slide, item)
        elif kind == "timeline":
            add_timeline_slide(slide, item)
        elif kind == "bmc":
            add_bmc_slide(slide)
        elif kind == "section":
            add_section_slide(slide, item)
        elif kind == "conclusion":
            add_conclusion_slide(slide, item)
        elif kind == "thanks":
            add_thanks_slide(slide, item)

    prs.save(OUT_PPTX)


def font(size: int, bold: bool = False, italic: bool = False):
    base = "/usr/share/fonts/truetype/dejavu"
    if bold and italic:
        name = "DejaVuSerif-Bold.ttf"
    elif bold:
        name = "DejaVuSans-Bold.ttf"
    elif italic:
        name = "DejaVuSerif.ttf"
    else:
        name = "DejaVuSans.ttf"
    return ImageFont.truetype(str(Path(base) / name), size=size)


def draw_wrapped(draw: ImageDraw.ImageDraw, text: str, xy, max_width: int, fnt, fill, spacing: int = 6, align: str = "left"):
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=fnt) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)

    x, y = xy
    for line in lines:
        if align == "center":
            lx = x + (max_width - draw.textlength(line, font=fnt)) / 2
        else:
            lx = x
        draw.text((lx, y), line, font=fnt, fill=fill)
        y += fnt.size + spacing
    return y


def render_pdf_slide(item: dict, idx: int) -> Image.Image:
    w, h = 1600, 900
    img = Image.new("RGB", (w, h), PIL_COLORS["white"])
    draw = ImageDraw.Draw(img)

    # Decoration.
    for i in range(5):
        draw.ellipse((1360 + i * 30, 45, 1376 + i * 30, 61), fill=PIL_COLORS["navy"])
        draw.ellipse((70 + i * 27, 835, 86 + i * 27, 851), fill=PIL_COLORS["navy"])
    draw.rectangle((70, 88, 500, 92), fill=PIL_COLORS["muted"])
    draw.rectangle((1080, 815, 1460, 819), fill=PIL_COLORS["muted"])
    draw.text((70, 25), item["section"], font=font(13, bold=True), fill=PIL_COLORS["muted"])
    draw.text((1460, 848), str(idx), font=font(13), fill=PIL_COLORS["muted"])

    if item["type"] == "title":
        draw.text((w / 2, 240), item["title"], font=font(76, bold=True, italic=True), fill=PIL_COLORS["navy"], anchor="mm")
        draw.text((w / 2, 330), item["subtitle"], font=font(35), fill=PIL_COLORS["navy"], anchor="mm")
        y = 405
        for line in item["body"]:
            draw.text((w / 2, y), line, font=font(22), fill=PIL_COLORS["muted"], anchor="mm")
            y += 42
        return img

    if item["type"] == "section":
        draw.text((w / 2, 360), item["title"], font=font(52, bold=True, italic=True), fill=PIL_COLORS["navy"], anchor="mm")
        draw.text((w / 2, 430), item["subtitle"], font=font(28), fill=PIL_COLORS["muted"], anchor="mm")
        draw.rectangle((560, 495, 1040, 503), fill=PIL_COLORS["teal"])
        return img

    if item["type"] == "thanks":
        draw.text((w / 2, 380), item["title"], font=font(48, bold=True, italic=True), fill=PIL_COLORS["navy"], anchor="mm")
        draw.text((w / 2, 455), item["subtitle"], font=font(34), fill=PIL_COLORS["teal"], anchor="mm")
        return img

    draw.text((75, 75), item["title"], font=font(36, bold=True, italic=True), fill=PIL_COLORS["navy"])
    if item.get("subtitle"):
        draw.text((80, 132), item["subtitle"], font=font(19), fill=PIL_COLORS["muted"])

    if item["type"] == "sommaire":
        y = 210
        for i, entry in enumerate(item["items"], start=1):
            x = 135 if i <= 5 else 850
            yy = y + ((i - 1) % 5) * 72
            draw.rounded_rectangle((x - 35, yy - 12, x + 560, yy + 45), radius=18, fill=PIL_COLORS["light"])
            draw.text((x, yy), f"{i}. {entry}", font=font(20), fill=PIL_COLORS["navy_dark"])
    elif item["type"] in {"cards", "bmc"}:
        if item["type"] == "bmc":
            card_data = [
                ("Partenaires clés", ["Microsoft Azure", "Open source", "Écoles / MSSP"], "blue"),
                ("Activités clés", ["Terraform", "KQL / SOAR", "ML et support"], "teal"),
                ("Proposition de valeur", ["SOC cloud intelligent", "automatisé et économique"], "orange"),
                ("Relations clients", ["Documentation", "Onboarding", "Support"], "blue"),
                ("Segments clients", ["PME, startups", "RSSI, MSSP", "Écoles"], "teal"),
                ("Ressources clés", ["Azure, Sentinel", "SQL, Logic Apps", "Scripts Python"], "orange"),
                ("Canaux", ["GitHub", "Azure Marketplace", "LinkedIn"], "blue"),
                ("Structure de coûts", ["Infrastructure Azure", "Maintenance", "Support / R&D"], "teal"),
                ("Sources de revenus", ["Freemium", "SaaS", "Enterprise"], "orange"),
            ]
        else:
            card_data = item["cards"]
        cols = 4 if len(card_data) == 4 else 3
        for i, (title, bullets, color_name) in enumerate(card_data):
            col = i % cols
            row = i // cols
            x = 90 + col * (1420 / cols)
            y = 210 + row * 185
            cw = 1420 / cols - 28
            ch = 155 if len(card_data) > 4 else 420
            draw.rounded_rectangle((x, y, x + cw, y + ch), radius=18, fill=PIL_COLORS["light"], outline=PIL_COLORS["muted"])
            draw.rectangle((x, y, x + 10, y + ch), fill=PIL_COLORS[color_name])
            draw.text((x + 28, y + 22), title, font=font(20, bold=True), fill=PIL_COLORS["navy"])
            by = y + 65
            for bullet in bullets:
                by = draw_wrapped(draw, f"- {bullet}", (x + 32, by), int(cw - 60), font(16), PIL_COLORS["navy_dark"], spacing=3)
    elif item["type"] in {"architecture_detail", "class_diagram", "sequence_diagram", "usecase_diagram"}:
        if item["type"] == "architecture_detail":
            zones = [
                ("VNet1 - DMZ", ["VM1 WebServer", "NSG-DMZ", "Public IP"], "blue", 90, 300),
                ("VNet2 - Data", ["Azure SQL", "Blob Storage", "Data Factory"], "teal", 560, 300),
                ("VNet3 - SIEM/SOAR", ["Sentinel", "Logic Apps", "VM2 ML"], "navy", 1030, 300),
            ]
            draw.rounded_rectangle((620, 165, 975, 245), radius=18, fill=PIL_COLORS["red"])
            draw.text((797, 190), "Internet / Attaquant", font=font(20, bold=True), fill=PIL_COLORS["white"], anchor="ma")
            for title, comps, color_name, x, y in zones:
                draw.rounded_rectangle((x, y, x + 380, y + 300), radius=20, outline=PIL_COLORS[color_name], width=3)
                draw.text((x + 190, y + 25), title, font=font(19, bold=True), fill=PIL_COLORS[color_name], anchor="ma")
                for i, comp in enumerate(comps):
                    yy = y + 70 + i * 68
                    draw.rounded_rectangle((x + 45, yy, x + 335, yy + 48), radius=12, fill=PIL_COLORS[color_name])
                    draw.text((x + 190, yy + 14), comp, font=font(16, bold=True), fill=PIL_COLORS["white"], anchor="ma")
            draw.rounded_rectangle((600, 645, 950, 700), radius=15, fill=PIL_COLORS["black"])
            draw.text((775, 662), "Log Analytics Workspace", font=font(17, bold=True), fill=PIL_COLORS["white"], anchor="ma")
            draw.rounded_rectangle((1080, 645, 1395, 700), radius=15, fill=PIL_COLORS["orange"])
            draw.text((1237, 662), "Power BI", font=font(18, bold=True), fill=PIL_COLORS["white"], anchor="ma")
        elif item["type"] == "class_diagram":
            names = ["AttackEvent", "HoneypotVM", "LogPipeline", "MLModel", "SIEMAlert", "SOARPlaybook", "NSGRule", "PowerBIDashboard"]
            colors = ["red", "blue", "navy", "navy", "teal", "teal", "orange", "black"]
            for i, name in enumerate(names):
                col, row = i % 4, i // 4
                x, y = 110 + col * 370, 220 + row * 230
                draw.rounded_rectangle((x, y, x + 300, y + 170), radius=18, fill=PIL_COLORS[colors[i]])
                draw.text((x + 150, y + 20), name, font=font(19, bold=True), fill=PIL_COLORS["white"], anchor="ma")
                draw.line((x + 20, y + 58, x + 280, y + 58), fill=PIL_COLORS["white"])
                draw_wrapped(draw, "- attributs principaux\n+ méthodes principales", (x + 28, y + 78), 245, font(15), PIL_COLORS["white"], spacing=6)
            draw.text((800, 725), "Chaîne logique : événement -> pipeline -> alerte -> réponse -> visualisation", font=font(18), fill=PIL_COLORS["muted"], anchor="ma")
        elif item["type"] == "sequence_diagram":
            actors = ["Attaquant", "DVWA", "Suricata", "Pipeline", "Sentinel", "Logic Apps", "NSG"]
            xs = [150 + i * 210 for i in range(len(actors))]
            for x, actor in zip(xs, actors):
                draw.rounded_rectangle((x - 70, 185, x + 70, 235), radius=12, fill=PIL_COLORS["navy"])
                draw.text((x, 200), actor, font=font(15, bold=True), fill=PIL_COLORS["white"], anchor="ma")
                draw.line((x, 250, x, 690), fill=PIL_COLORS["light"], width=3)
            steps = [
                (0, 1, "1. Attaque HTTP/SSH", 285),
                (1, 2, "2. EVE JSON", 340),
                (1, 3, "3. Logs Apache/auth", 395),
                (3, 4, "4. Alerte SIEM", 470),
                (4, 5, "5. Déclenche SOAR", 540),
                (5, 6, "6. Bloquer IP", 610),
            ]
            for src, dst, label, y in steps:
                draw.line((xs[src], y, xs[dst], y), fill=PIL_COLORS["teal"], width=4)
                draw.polygon([(xs[dst], y), (xs[dst] - 12, y - 7), (xs[dst] - 12, y + 7)], fill=PIL_COLORS["teal"])
                draw.text(((xs[src] + xs[dst]) / 2, y - 25), label, font=font(15), fill=PIL_COLORS["navy_dark"], anchor="ma")
        elif item["type"] == "usecase_diagram":
            draw.rounded_rectangle((355, 165, 1245, 695), radius=24, outline=PIL_COLORS["light"], width=4)
            draw.text((800, 190), "Plateforme AITDR", font=font(20, bold=True), fill=PIL_COLORS["muted"], anchor="ma")
            usecases = [
                ("Capturer des attaques", 450, 260, "blue"),
                ("Collecter les logs", 450, 350, "blue"),
                ("Détecter anomalies (ML)", 450, 440, "blue"),
                ("Générer alertes SIEM", 450, 530, "teal"),
                ("Réponse auto. SOAR", 450, 620, "teal"),
                ("Visualiser tableau bord", 840, 305, "navy"),
                ("Consulter incidents", 840, 405, "navy"),
                ("Configurer règles KQL", 840, 505, "navy"),
                ("Déployer infrastructure", 840, 605, "orange"),
            ]
            for text, x, y, color_name in usecases:
                draw.ellipse((x, y, x + 320, y + 58), fill=PIL_COLORS[color_name])
                draw.text((x + 160, y + 18), text, font=font(16, bold=True), fill=PIL_COLORS["white"], anchor="ma")
            draw.text((160, 370), "Attaquant", font=font(19), fill=PIL_COLORS["muted"], anchor="ma")
            draw.text((165, 610), "Analyste SOC", font=font(19), fill=PIL_COLORS["muted"], anchor="ma")
            draw.text((1415, 480), "Admin", font=font(19), fill=PIL_COLORS["muted"], anchor="ma")
    else:
        # Compact fallback for PDF rendering.
        y = 215
        if item.get("paragraph"):
            y = draw_wrapped(draw, item["paragraph"], (100, y), 1380, font(23), PIL_COLORS["navy_dark"], spacing=8)
            y += 25
        bullets = item.get("bullets", [])
        if item.get("metrics"):
            for i, (value, label, color_name) in enumerate(item["metrics"]):
                x = 115 + i * 365
                draw.rounded_rectangle((x, 205, x + 310, 330), radius=18, fill=PIL_COLORS["light"])
                draw.text((x + 155, 240), value, font=font(27, bold=True), fill=PIL_COLORS[color_name], anchor="mm")
                draw.text((x + 155, 292), label, font=font(15), fill=PIL_COLORS["navy_dark"], anchor="mm")
            y = 405
        if item.get("columns"):
            y = 205
            draw.rounded_rectangle((95, y, 1505, y + 500), radius=18, fill=PIL_COLORS["light"])
            rows = [item["columns"]] + item["rows"]
            row_h = 500 / len(rows)
            col_w = 1410 / len(rows[0])
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    x = 95 + c * col_w
                    yy = y + r * row_h
                    fill = PIL_COLORS["navy"] if r == 0 else ((255, 255, 255) if r % 2 else PIL_COLORS["light"])
                    draw.rectangle((x, yy, x + col_w, yy + row_h), fill=fill, outline=PIL_COLORS["muted"])
                    draw_wrapped(draw, val, (x + 10, yy + 12), int(col_w - 20), font(14, bold=r == 0), PIL_COLORS["white"] if r == 0 else PIL_COLORS["navy_dark"], spacing=2)
            bullets = []
        if item.get("nodes") or item.get("steps"):
            elements = item.get("nodes") or item.get("steps")
            for i, node in enumerate(elements):
                x = 115 + i * (1370 / len(elements))
                draw.rounded_rectangle((x, 300, x + 230, 430), radius=18, fill=PIL_COLORS["light"], outline=PIL_COLORS["muted"])
                draw.text((x + 115, 325), node[0 if item.get("nodes") else 1], font=font(18, bold=True), fill=PIL_COLORS["navy"], anchor="ma")
                draw_wrapped(draw, node[1 if item.get("nodes") else 2], (x + 20, 355), 190, font(14), PIL_COLORS["navy_dark"], align="center")
            bullets = []
        for bullet in bullets:
            y = draw_wrapped(draw, f"- {bullet}", (120, y), 1360, font(22), PIL_COLORS["navy_dark"], spacing=6)

    return img


def build_pdf():
    images = [render_pdf_slide(item, idx) for idx, item in enumerate(SLIDES, start=1)]
    first, rest = images[0], images[1:]
    first.save(OUT_PDF, save_all=True, append_images=rest, resolution=144.0)
    first.save(OUT_PDF_COMPAT, save_all=True, append_images=rest, resolution=144.0)


def build_notes():
    lines = [
        "# Notes de présentation - Partie Conception AITDR",
        "",
        "Ces notes sont prévues pour être collées dans les notes du présentateur.",
        "",
    ]
    for idx, item in enumerate(SLIDES, start=1):
        if "talk" not in item:
            continue
        lines.extend(
            [
                f"## Slide {idx} - {item['title']}",
                "",
                item["talk"],
                "",
            ]
        )
    OUT_NOTES.write_text("\n".join(lines), encoding="utf-8")


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_pptx()
    build_pdf()
    build_notes()
    print(f"Generated {OUT_PPTX}")
    print(f"Generated {OUT_PDF}")
    print(f"Generated {OUT_PDF_COMPAT}")
    print(f"Generated {OUT_NOTES}")
    print(f"Slides: {len(SLIDES)}")


if __name__ == "__main__":
    main()
